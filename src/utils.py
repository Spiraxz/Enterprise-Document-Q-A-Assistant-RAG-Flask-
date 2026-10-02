from langchain_google_genai import GoogleGenerativeAIEmbeddings, ChatGoogleGenerativeAI
from langchain_community.vectorstores.chroma import Chroma
from langchain_text_splitters import RecursiveCharacterTextSplitter
from src.logger import logging
from dotenv.main import load_dotenv
import os
import io
load_dotenv()


SUPPORTED_EXTENSIONS=('.txt', '.md', '.pdf', '.docx')

def extract_text(file_storage, ext):
    """Extract plain text from an uploaded file based on its extension."""
    data=file_storage.read()
    if ext in ('.txt', '.md'):
        return data.decode('utf-8', errors='ignore')
    if ext == '.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or '' for page in reader.pages)
    if ext == '.docx':
        import docx
        document=docx.Document(io.BytesIO(data))
        parts=[p.text for p in document.paragraphs]
        parts+= [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join(parts)
    raise ValueError(f"Unsupported file type '{ext}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}")


google_api_key=os.getenv('GOOGLE_API_KEY') or os.getenv('GEMINI_API_KEY')
GEMINI_MODEL=os.getenv('GEMINI_MODEL') or 'gemini-3.5-flash-lite'
class Utils:

    def __init__(self):
        self.model=GoogleGenerativeAIEmbeddings(
        google_api_key=google_api_key,
        model='gemini-embedding-2-preview')
        self.llm=ChatGoogleGenerativeAI(model=GEMINI_MODEL,google_api_key=google_api_key)
        self.vectorStore=Chroma('langchain_store',self.model,persist_directory='./database')
        self.vectorStore.persist()
        
        
        
    def add_text(self,input):
        self.vectorStore.add_texts([input])

    def add_document(self,text):
        """Chunk a document and store its chunks in the vector database.
        Returns the number of chunks added."""
        splitter=RecursiveCharacterTextSplitter(chunk_size=500,chunk_overlap=50)
        chunks=splitter.split_text(text)
        if chunks:
            self.vectorStore.add_texts(chunks)
        return len(chunks)

    def findmatch(self,input):
        result=self.vectorStore.similarity_search(query=input)
        if result==[]:
            return result
        else:
            response=""
            for i in range (len(result)):
                response+=result[i].page_content+" "
            return response

    def query_refiner(self,conversation,query):
        response=self.llm.invoke(
            f"""###SYSTEM: You are a search-query rewriter for a document Q&A assistant. Using the conversation so far, rewrite the user's query as a concise standalone search query that captures their intent (resolve words like 'it', 'that', 'them' using the conversation). If there is no conversation, return the query unchanged. Output only the rewritten search query, nothing else.
            ###TEXT: {conversation} 
            ###USER:{query}?
            ###RESPONSE: """
        )
        refined=str(response.text).strip().strip('"')
        return refined or query

    def get_conversation_string(self,requests, responses):
        conversation_string=""
        for i in range(len(responses)-1):
            conversation_string+=f"User: {requests[i]}\n"
            conversation_string+=f"Bot: {responses[i+1]}\n"
        logging.info(conversation_string)
        return conversation_string    
    
    def get_all_docs(self):
        db=self.vectorStore.get()
        conversation=db['documents']
        return conversation
